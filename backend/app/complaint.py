"""Downloadable complaint drafts (.docx) built from the analysis.

Three drafts, matching the escalation path:
  fir         - written complaint to the officer in charge of the police station, asking to register an FIR
  sp          - written complaint to the Superintendent of Police when the station refuses
  magistrate  - application to the Magistrate to direct an investigation (usually filed with a lawyer)

Procedural sections are taken from the LegalCompass mapping for the code that governs the proceedings (BNSS or
CrPC). Everything the user must fill in is in [square brackets] and highlighted. The user's own words are kept
as written (this is their complaint, so nothing is masked).
"""
from __future__ import annotations

import io
import re
from datetime import date

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Pt

from . import lc_client

ACTS = {"BNSS": "the Bharatiya Nagarik Suraksha Sanhita, 2023", "CrPC": "the Code of Criminal Procedure, 1973"}
KINDS = {"fir": "Complaint to Police Station (FIR)", "sp": "Complaint to Superintendent of Police",
         "magistrate": "Application to Magistrate"}


def _proc_section(crpc: str, procedure: str) -> str:
    """'Section 173 of the Bharatiya Nagarik Suraksha Sanhita, 2023' or the CrPC original."""
    if procedure == "BNSS":
        try:
            r = lc_client.lookup("CrPC", crpc) or {}
            secs = sorted({x["new"]["section"] for x in r.get("results", [])})
            if secs:
                return f"Section {' / '.join(secs)} of {ACTS['BNSS']}"
        except lc_client.LegalCompassUnavailable:
            pass
    return f"Section {crpc} of {ACTS['CrPC']}"


def _sections_line(analysis: dict) -> list[str]:
    law = analysis["regime"].get("law")
    out = []
    for s in analysis["sections"]:
        bns = " / ".join(b["provision"] for b in s["bns"])
        if law == "BNS" and bns:
            out.append(f"{bns} (formerly IPC {s['ipc']}) — {s['ipc_heading'].rstrip('.')}")
        elif law == "IPC":
            out.append(f"IPC {s['ipc']} — {s['ipc_heading'].rstrip('.')}")
        else:
            out.append(f"IPC {s['ipc']} / {bns or 'no BNS counterpart'} — {s['ipc_heading'].rstrip('.')}")
    return out


def _para(doc, text: str = "", bold: bool = False, size: int = 11, align=None, space_after: int = 6):
    """Add a paragraph; [placeholders] are highlighted so the user sees what to fill in."""
    p = doc.add_paragraph()
    for part in re.split(r"(\[[^\]]+\])", text):
        if not part:
            continue
        run = p.add_run(part)
        run.bold = bold
        run.font.size = Pt(size)
        if part.startswith("[") and part.endswith("]"):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def build(kind: str, narrative: str, analysis: dict) -> tuple[bytes, str]:
    if kind not in KINDS:
        raise ValueError("kind must be fir, sp or magistrate")
    a = analysis["answers"]
    procedure = analysis.get("procedure_code") or "BNSS"
    secs = _sections_line(analysis)
    s154, s156 = _proc_section("154", procedure), _proc_section("156", procedure)
    try:
        od = date.fromisoformat(a["offence_date"]).strftime("%d %B %Y")
    except (KeyError, ValueError, TypeError):
        od = "[date of incident]"
    state = a.get("state") or "[State]"
    doc = Document()
    for st in doc.styles:
        if hasattr(st, "font") and st.font is not None:
            st.font.name = "Calibri"

    note = doc.add_paragraph()
    r = note.add_run("DRAFT prepared with LegalCompass — check every detail, fill in the highlighted [brackets], "
                     "and have it reviewed by a lawyer or Legal Services Authority (15100) if possible. Not legal advice.")
    r.italic = True
    r.font.size = Pt(9)

    if kind == "magistrate":
        _para(doc, "IN THE COURT OF THE [Judicial Magistrate First Class / Metropolitan Magistrate], [Place]",
              bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _para(doc, f"Application under {s156}", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        _para(doc, "[Your full name], [son / daughter / wife] of [name], aged [age], residing at [full address]  … Applicant")
        _para(doc, "Versus")
        _para(doc, "[Name(s) of accused, if known, with address] / Unknown persons  … Accused")
        _para(doc, "")
        _para(doc, "The Applicant respectfully submits:", bold=True)
    else:
        _para(doc, f"Date: {date.today().strftime('%d %B %Y')}")
        _para(doc, "To,")
        if kind == "fir":
            _para(doc, f"The Officer in Charge / Station House Officer,\n[Name] Police Station,\n[District], {state}")
            subject = f"Request to register an FIR under {s154}"
        else:
            _para(doc, f"The Superintendent of Police / Deputy Commissioner of Police,\n[District / Zone], {state}")
            subject = (f"Complaint under {s154} — the officer in charge of [Name] Police Station did not register my "
                       "information as an FIR")
        _para(doc, f"Subject: {subject}", bold=True)
        _para(doc, "Respected Sir / Madam,")
        _para(doc, "I, [Your full name], [son / daughter / wife] of [name], aged [age], residing at [full address], "
                   "phone [mobile number], wish to report the following:")

    n = 1
    _para(doc, f"{n}. On {od}, at about [time], at [place of incident], the following happened:", bold=True); n += 1
    for line in [l.strip() for l in narrative.strip().splitlines() if l.strip()]:
        _para(doc, line)
    facts = [("Accused", "[name(s), description, address if known]"),
             ("Number of persons involved", a.get("accused_count")),
             ("Injury", a.get("injury")), ("Money / property involved", a.get("amount")),
             ("Witnesses", "[names and phone numbers, if any]")]
    _para(doc, f"{n}. Details:", bold=True); n += 1
    for k, v in facts:
        if v:
            _para(doc, f"    • {k}: {v}")
    if secs:
        _para(doc, f"{n}. The above facts appear to disclose offences that may fall under:", bold=True); n += 1
        for s in secs:
            _para(doc, f"    • {s}")
        _para(doc, "    (The exact sections are for the police / court to decide.)", size=9)
    if kind == "sp":
        _para(doc, f"{n}. On [date], I approached [Name] Police Station to report this, but the officer in charge "
                   "[refused to register an FIR / did not take any action]. [Give the officer's name if known and "
                   "what was said.]", bold=False); n += 1
    if kind == "magistrate":
        _para(doc, f"{n}. On [date] the Applicant approached [Name] Police Station, and on [date] sent a written "
                   f"complaint to the Superintendent of Police under {s154}. No FIR has been registered and no "
                   "investigation has begun. [Attach copies and postal / email receipts.]"); n += 1
        _para(doc, f"{n}. The facts disclose a cognizable offence which requires investigation by the police.");  n += 1
        _para(doc, "PRAYER", bold=True)
        _para(doc, "It is therefore prayed that this Hon'ble Court may be pleased to direct the officer in charge of "
                   "[Name] Police Station to register an FIR and investigate the matter, and pass any other order "
                   "it deems fit.")
        _para(doc, "This application is supported by an affidavit of the Applicant.", size=10)
        _para(doc, "")
        _para(doc, "Place: [Place]\nDate: [Date]")
        _para(doc, "Signature of Applicant: ____________________\n[Your full name]\n(through Advocate [name], if any)")
    else:
        if kind == "fir":
            _para(doc, f"{n}. I request you to register an FIR on the basis of this information, investigate the matter "
                       "and take action against the accused in accordance with law, and give me a free copy of the FIR.")
        else:
            _para(doc, f"{n}. I request you to look into this complaint, and either investigate it yourself or direct "
                       "an officer to register an FIR and investigate, in accordance with law.")
        n += 1
        _para(doc, "I declare that the information given above is true to the best of my knowledge.")
        _para(doc, "")
        _para(doc, "Yours faithfully,\n\nSignature: ____________________\nName: [Your full name]\nPhone: [mobile number]")
    _para(doc, "Enclosures: [copies of ID, medical report, bills / bank statements, screenshots, earlier complaint "
               "and receipt — as applicable]", size=10)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), f"legalcompass-{kind}-draft.docx"
